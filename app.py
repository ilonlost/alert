import os
import sys
import smtplib
import traceback
from datetime import datetime, time
from collections import Counter
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
from zoneinfo import ZoneInfo

import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from apscheduler.schedulers.blocking import BlockingScheduler
from dotenv import load_dotenv

load_dotenv()

DB_SERVER = os.getenv("DB_SERVER", "11-vm-dwh01")
DB_NAME = os.getenv("DB_NAME", "TimeShiftFK")
DB_USER = os.getenv("DB_USER", "")
DB_PASSWORD = os.getenv("DB_PASSWORD", "")
DB_DRIVER = os.getenv("DB_DRIVER", "FreeTDS")
DB_PORT = os.getenv("DB_PORT", "1433")
DB_TDS_VERSION = os.getenv("DB_TDS_VERSION", "7.4")
DB_TRUST_CERT = os.getenv("DB_TRUST_CERT", "yes")

SMTP_HOST = os.getenv("SMTP_HOST", "smtp.agrohold.ru")
SMTP_PORT = int(os.getenv("SMTP_PORT", "25"))
SMTP_USER = os.getenv("SMTP_USER", "").strip()
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "")
MAIL_FROM = os.getenv("MAIL_FROM", "WGSentry11@agrohold.ru")
MAIL_TO = [x.strip() for x in os.getenv("MAIL_TO", "").split(",") if x.strip()]

TARGET_AREA = os.getenv("TARGET_AREA", "ФК-1эт-ВЕСТИБЮЛЬ(ЛК9)")
TIMEZONE = os.getenv("TIMEZONE", "Europe/Moscow")
TZ = ZoneInfo(TIMEZONE)

MORNING_START = time(6, 45)
MORNING_END = time(8, 20)
EVENING_START = time(18, 45)
EVENING_END = time(20, 20)

REPORT_DIR = os.getenv("REPORT_DIR", "/app/reports")
os.makedirs(REPORT_DIR, exist_ok=True)


def log(message):
    print(f"[{datetime.now(TZ):%Y-%m-%d %H:%M:%S}] {message}", flush=True)


def normalize_code(value):
    return "" if value is None else str(value).strip().upper()


def row_to_dict(columns, row):
    return {str(k).lower(): v for k, v in zip(columns, row)}


def get_connection():
    drivers = pyodbc.drivers()
    driver = DB_DRIVER
    if driver not in drivers:
        if "FreeTDS" in drivers:
            driver = "FreeTDS"
        elif "ODBC Driver 18 for SQL Server" in drivers:
            driver = "ODBC Driver 18 for SQL Server"
        elif "ODBC Driver 17 for SQL Server" in drivers:
            driver = "ODBC Driver 17 for SQL Server"
        else:
            raise RuntimeError(f"No SQL Server ODBC driver found. Installed: {drivers}")

    conn_str = (
        f"DRIVER={{{driver}}};"
        f"SERVER={DB_SERVER};"
        f"DATABASE={DB_NAME};"
        f"UID={DB_USER};"
        f"PWD={DB_PASSWORD};"
    )
    if driver == "FreeTDS":
        conn_str += f"PORT={DB_PORT};TDS_Version={DB_TDS_VERSION};ClientCharset=UTF-8;"
    else:
        conn_str += f"TrustServerCertificate={DB_TRUST_CERT};"

    return pyodbc.connect(conn_str, autocommit=True)


def get_personnel(cursor):
    cursor.execute("EXEC dbo.GetFKPersonal")
    personnel = {}

    while True:
        if cursor.description:
            columns = [c[0] for c in cursor.description]
            rows = cursor.fetchall()
            lower_columns = [str(c).lower() for c in columns]

            if "code" in lower_columns:
                for row in rows:
                    data = row_to_dict(columns, row)
                    code = normalize_code(data.get("code"))
                    if not code:
                        continue

                    fio = " ".join(
                        x for x in [
                            str(data.get("lastname") or "").strip(),
                            str(data.get("firstname") or "").strip(),
                            str(data.get("middlename") or "").strip(),
                        ] if x
                    )

                    personnel[code] = {
                        "fio": fio,
                        "short_department": str(data.get("staffroutinedivisions") or "").strip(),
                        "full_department": str(data.get("routinedivisionsfullname") or "").strip(),
                        "position": str(data.get("staffroutinepositions") or "").strip(),
                    }

        if not cursor.nextset():
            break

    log(f"Personnel loaded: {len(personnel)}")
    return personnel


def get_movements(cursor, date_start, date_end):
    cursor.execute(
        "EXEC dbo.GetFKPersonalMovements @dateStart = ?, @dateEnd = ?",
        date_start,
        date_end,
    )

    movements = []
    while True:
        if cursor.description:
            columns = [c[0] for c in cursor.description]
            lower_columns = [str(c).lower() for c in columns]
            rows = cursor.fetchall()

            if "code" in lower_columns and "repdate" in lower_columns:
                for row in rows:
                    movements.append(row_to_dict(columns, row))

        if not cursor.nextset():
            break

    log(f"Movement rows loaded: {len(movements)}")
    return movements


def is_target_area(movement):
    target = TARGET_AREA.lower()
    area = str(movement.get("areaname") or "").strip().lower()
    terminal = str(movement.get("terminalname") or "").strip().lower()
    return target in area or target in terminal


def extract_department(employee):
    full_department = str(employee.get("full_department") or "").strip()
    if full_department:
        parts = [p.strip() for p in full_department.split("/") if p.strip()]
        if len(parts) >= 2:
            return parts[1]
        if parts:
            return parts[0]

    short_department = str(employee.get("short_department") or "").strip()
    return short_department or "Не определено"


def first_entries(movements):
    result = {}
    for movement in movements:
        code = normalize_code(movement.get("code"))
        if not code:
            continue

        if str(movement.get("typepass") or "").strip().upper() != "ВХОД":
            continue
        if not is_target_area(movement):
            continue

        rep_date = movement.get("repdate")
        if not rep_date:
            continue

        if code not in result or rep_date < result[code]["repdate"]:
            result[code] = movement

    return result


def build_report(report_type):
    now = datetime.now(TZ)
    today = now.date()

    if report_type == "morning":
        start_t, end_t, title = MORNING_START, MORNING_END, "Утренний отчёт"
    elif report_type == "evening":
        start_t, end_t, title = EVENING_START, EVENING_END, "Вечерний отчёт"
    else:
        raise ValueError("report_type must be morning or evening")

    date_start = datetime.combine(today, start_t).replace(tzinfo=None)
    date_end = datetime.combine(today, end_t).replace(tzinfo=None)

    log(f"Building {report_type} report for {date_start:%H:%M}-{date_end:%H:%M}")

    with get_connection() as conn:
        cursor = conn.cursor()
        personnel = get_personnel(cursor)
        movements = get_movements(cursor, date_start, date_end)

    entries = first_entries(movements)
    departments = Counter()
    employees = []

    for code, movement in entries.items():
        employee = personnel.get(code)

        if employee:
            department = extract_department(employee)
            fio = employee["fio"]
            full_department = employee["full_department"]
            position = employee["position"]
        else:
            # This keeps the sum of departments equal to "Всего прошло".
            department = "Не найден в справочнике"
            fio = ""
            full_department = ""
            position = ""

        departments[department] += 1
        employees.append({
            "code": code,
            "fio": fio,
            "time": movement["repdate"],
            "department": department,
            "full_department": full_department,
            "position": position,
        })

    department_counts = sorted(departments.items(), key=lambda x: (-x[1], x[0].lower()))
    total = len(entries)
    department_sum = sum(count for _, count in department_counts)

    if total != department_sum:
        raise RuntimeError(f"Integrity check failed: total={total}, department_sum={department_sum}")

    report = {
        "type": report_type,
        "title": title,
        "date": today,
        "start": date_start,
        "end": date_end,
        "total": total,
        "department_counts": department_counts,
        "employees": employees,
    }

    log(f"Unique employees: {total}; department sum: {department_sum}")
    for dept, count in department_counts:
        log(f"  {dept}: {count}")

    return report


def create_excel(report):
    filename = f"access_{report['type']}_{report['date']:%Y-%m-%d}.xlsx"
    path = os.path.join(REPORT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Проходы"

    headers = [
        "№", "Код сотрудника", "ФИО", "Время прохода", "Подразделение",
        "Полный путь подразделения", "Должность", "СКУД"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, employee in enumerate(sorted(report["employees"], key=lambda x: x["time"]), 1):
        ws.append([
            i,
            employee["code"],
            employee["fio"],
            employee["time"].strftime("%H:%M:%S"),
            employee["department"],
            employee["full_department"],
            employee["position"],
            TARGET_AREA,
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"A": 7, "B": 20, "C": 38, "D": 18, "E": 35, "F": 70, "G": 40, "H": 35}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    stats = wb.create_sheet("Статистика")
    stats.append(["Показатель", "Количество"])
    stats["A1"].font = Font(bold=True)
    stats["B1"].font = Font(bold=True)
    stats.append(["Всего прошло", report["total"]])
    stats.append([])
    stats.append(["Подразделение", "Количество"])
    stats["A4"].font = Font(bold=True)
    stats["B4"].font = Font(bold=True)
    for department, count in report["department_counts"]:
        stats.append([department, count])
    stats.append([])
    stats.append(["Сумма по подразделениям", sum(x[1] for x in report["department_counts"])])
    stats.column_dimensions["A"].width = 50
    stats.column_dimensions["B"].width = 22

    wb.save(path)
    log(f"Excel created: {path}")
    return path


def department_html(departments):
    return "".join(
        f"<tr><td>{name}</td><td style='text-align:center'><b>{count}</b></td></tr>"
        for name, count in departments
    ) or "<tr><td colspan='2'>Нет данных</td></tr>"


def send_email(report, excel_path):
    if not MAIL_TO:
        raise RuntimeError("MAIL_TO is empty")

    subject = f"{report['title']} по проходам сотрудников {report['date']:%d.%m.%Y}"
    body = f"""
    <html><body style="font-family:Arial,Helvetica,sans-serif;color:#222;font-size:15px">
      <h2>{report['title']} по проходам сотрудников</h2>
      <p><b>Дата:</b> {report['date']:%d.%m.%Y}</p>
      <p><b>Период:</b> {report['start']:%H:%M} — {report['end']:%H:%M}</p>
      <p><b>Точка прохода:</b> {TARGET_AREA}</p>
      <table border="1" cellspacing="0" cellpadding="12" style="border-collapse:collapse;font-size:16px">
        <tr><td><b>Всего прошло</b></td><td style="text-align:center"><b>{report['total']}</b></td></tr>
      </table>
      <br>
      <h3>По подразделениям</h3>
      <table border="1" cellspacing="0" cellpadding="9" style="border-collapse:collapse">
        <tr><th>Подразделение</th><th>Количество</th></tr>
        {department_html(report['department_counts'])}
      </table>
      <br>
      <p>Во вложении Excel-файл с ФИО, временем прохода, подразделением и должностью.</p>
    </body></html>
    """

    msg = MIMEMultipart()
    msg["Subject"] = subject
    msg["From"] = MAIL_FROM
    msg["To"] = ", ".join(MAIL_TO)
    msg.attach(MIMEText(body, "html", "utf-8"))

    with open(excel_path, "rb") as f:
        attachment = MIMEApplication(
            f.read(),
            _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    attachment.add_header("Content-Disposition", "attachment", filename=os.path.basename(excel_path))
    msg.attach(attachment)

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
        if SMTP_USER and SMTP_PASSWORD:
            smtp.login(SMTP_USER, SMTP_PASSWORD)
        smtp.sendmail(MAIL_FROM, MAIL_TO, msg.as_string())

    log(f"Email sent to: {', '.join(MAIL_TO)}")


def run_report(report_type, send=True):
    report = build_report(report_type)
    excel_path = create_excel(report)
    if send:
        send_email(report, excel_path)
    else:
        log("TEST MODE: email not sent")
    return report, excel_path


def scheduled_morning():
    try:
        run_report("morning", send=True)
    except Exception:
        log("Morning report failed")
        traceback.print_exc()


def scheduled_evening():
    try:
        run_report("evening", send=True)
    except Exception:
        log("Evening report failed")
        traceback.print_exc()


def run_daemon():
    scheduler = BlockingScheduler(timezone=TIMEZONE)
    scheduler.add_job(scheduled_morning, "cron", hour=8, minute=30, id="morning", replace_existing=True)
    scheduler.add_job(scheduled_evening, "cron", hour=20, minute=30, id="evening", replace_existing=True)
    log(f"Scheduler started ({TIMEZONE}): morning 08:30, evening 20:30")
    scheduler.start()


def usage():
    print("Usage:")
    print("  python app.py daemon          # automatic 08:30 / 20:30")
    print("  python app.py morning         # run morning and SEND email")
    print("  python app.py evening         # run evening and SEND email")
    print("  python app.py test-morning    # run morning WITHOUT email")
    print("  python app.py test-evening    # run evening WITHOUT email")


def main():
    command = sys.argv[1].lower() if len(sys.argv) > 1 else "daemon"

    if command == "daemon":
        run_daemon()
    elif command == "morning":
        run_report("morning", send=True)
    elif command == "evening":
        run_report("evening", send=True)
    elif command == "test-morning":
        run_report("morning", send=False)
    elif command == "test-evening":
        run_report("evening", send=False)
    else:
        usage()
        sys.exit(1)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        log(f"ERROR: {exc!r}")
        traceback.print_exc()
        sys.exit(1)
